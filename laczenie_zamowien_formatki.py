# -*- coding: utf-8 -*-
"""
Tessainator V6 — program do łączenia zamówień / formatek na podstawie Excela.

Co robi:
- użytkownik wybiera plik Excel z okna,
- program nie łączy się z bazą danych,
- sprawdza pozycje według ustawień FEFCO i orientacji wymiarów, daty dostawy,
  dopuszczalnej przycinki oraz minimum sztukowego/metrowego,
- zapisuje wynik do nowego pliku Excel obok pliku wejściowego.

Instalacja zależności, jeśli ich brakuje:
    pip install pandas openpyxl

Opcjonalnie dla plików .ods:
    pip install odfpy
"""

from __future__ import annotations

import json
import math
import sys
import tempfile
import traceback
from dataclasses import dataclass
from datetime import datetime
from itertools import combinations
from pathlib import Path
from typing import Iterable, Optional

try:
    import pandas as pd
except ImportError:
    print("Brakuje biblioteki pandas. Zainstaluj: pip install pandas openpyxl")
    sys.exit(1)

try:
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk
except ImportError:
    tk = None
    filedialog = None
    messagebox = None
    ttk = None


# =========================
# USTAWIENIA PROGRAMU
# =========================

MAX_DNI_ROZNICY = 16
MAX_PRZYCINKA_PROC = 10.0
MIN_SZTUKI = 90
MIN_M2 = 290.0

# OR = grupa przechodzi, jeśli ma minimum sztuk ALBO minimum m2.
# AND = grupa przechodzi tylko wtedy, gdy spełnia oba minima naraz.
MINIMUM_WARUNEK = "OR"  # "OR" albo "AND"

# area = sprawdza procent uciętego pola formatu,
# side = sprawdza procent przycinki na każdym boku,
# both = sprawdza jednocześnie area i side.
TRYB_PRZYCINKI = "area"  # "area", "side" albo "both"

# Pozwala obracać formatki o 90 stopni podczas szukania wspólnego formatu,
# czyli porównywać długość jednej formatki z szerokością drugiej i odwrotnie.
LACZ_DLUGOSC_Z_SZEROKOSCIA = False

# Pozwala umieszczać w jednej grupie zamówienia o różnych kodach FEFCO.
LACZ_ROZNE_FEFCO = False

# Alternatywne wymiary porównujemy tylko dla FEFCO 201. Jeśli są oba zestawy,
# program traktuje je jako dwa różne rodzaje formatek, a nie wybiera tylko drugiego.
UZYJ_ALTERNATYWNYCH_DLA_FEFCO_201 = True

# Jeśli użyto alternatywnych wymiarów FEFCO 201:
# - alternative_1 oznacza 2 formatki i mnoży łączną liczbę sztuk przez 2,
# - alternative_1 + alternative_2 oznaczają dwa rodzaje formatek po 2 sztuki;
#   ich łączny mnożnik wynosi 4.
#
# Każda formatka jest dalej traktowana jako osobna pozycja dopasowania. Dzięki
# temu wszystkie formatki jednego kartonu mogą trafić do różnych grup.
MNOZNIK_SZTUK_ALT_1 = 2
MNOZNIK_SZTUK_ALT_2 = 4

# Jak w Twoim SQL: FEFCO 300 i 301 liczone x2 w m2/sztukach efektywnych.
MNOZNIK_FEFC0_300_301 = 2

# Po zbudowaniu grup program sprawdza, czy podział dużej grupy na mniejsze,
# nadal poprawne grupy zmniejszy łączny odpad.
OPTYMALIZUJ_ODPAD = True
MIN_OSZCZEDNOSC_ODPADU_M2 = 0.01

CONFIG_FILENAME = "ustawienia_laczenia.json"


@dataclass(frozen=True)
class TargetCheck:
    valid: bool
    target_x: float = math.nan
    target_y: float = math.nan
    max_trim_pct: float = math.nan
    avg_trim_pct: float = math.nan
    reason: str = ""
    rotated_indices: frozenset[int] = frozenset()


def show_error(title: str, text: str) -> None:
    if messagebox is not None:
        messagebox.showerror(title, text)
    else:
        print(f"{title}: {text}")


def show_info(title: str, text: str) -> None:
    if messagebox is not None:
        messagebox.showinfo(title, text)
    else:
        print(f"{title}: {text}")


def ustawienia_domyslne() -> dict:
    return {
        "MAX_DNI_ROZNICY": MAX_DNI_ROZNICY,
        "MAX_PRZYCINKA_PROC": MAX_PRZYCINKA_PROC,
        "MIN_SZTUKI": MIN_SZTUKI,
        "MIN_M2": MIN_M2,
        "MINIMUM_WARUNEK": MINIMUM_WARUNEK,
        "TRYB_PRZYCINKI": TRYB_PRZYCINKI,
        "LACZ_DLUGOSC_Z_SZEROKOSCIA": LACZ_DLUGOSC_Z_SZEROKOSCIA,
        "LACZ_ROZNE_FEFCO": LACZ_ROZNE_FEFCO,
        "UZYJ_ALTERNATYWNYCH_DLA_FEFCO_201": UZYJ_ALTERNATYWNYCH_DLA_FEFCO_201,
        "MNOZNIK_SZTUK_ALT_1": MNOZNIK_SZTUK_ALT_1,
        "MNOZNIK_SZTUK_ALT_2": MNOZNIK_SZTUK_ALT_2,
        "MNOZNIK_FEFC0_300_301": MNOZNIK_FEFC0_300_301,
        "OPTYMALIZUJ_ODPAD": OPTYMALIZUJ_ODPAD,
        "MIN_OSZCZEDNOSC_ODPADU_M2": MIN_OSZCZEDNOSC_ODPADU_M2,
    }


def sciezka_konfiguracji() -> Path:
    try:
        base_dir = Path(__file__).resolve().parent
    except Exception:
        base_dir = Path.cwd()
    return base_dir / CONFIG_FILENAME


def wczytaj_ustawienia_z_pliku() -> dict:
    settings = ustawienia_domyslne()
    path = sciezka_konfiguracji()
    if not path.exists():
        return settings
    try:
        loaded = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(loaded, dict):
            settings.update({k: loaded[k] for k in settings.keys() if k in loaded})
    except Exception:
        # Uszkodzony plik ustawień nie blokuje programu — używamy wartości domyślnych.
        return ustawienia_domyslne()
    return settings


def zapisz_ustawienia_do_pliku(settings: dict) -> None:
    try:
        sciezka_konfiguracji().write_text(
            json.dumps(settings, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
    except Exception:
        # Zapis ustawień jest wygodą, nie warunkiem działania programu.
        pass


def zastosuj_ustawienia(settings: dict) -> None:
    global MAX_DNI_ROZNICY
    global MAX_PRZYCINKA_PROC
    global MIN_SZTUKI
    global MIN_M2
    global MINIMUM_WARUNEK
    global TRYB_PRZYCINKI
    global LACZ_DLUGOSC_Z_SZEROKOSCIA
    global LACZ_ROZNE_FEFCO
    global UZYJ_ALTERNATYWNYCH_DLA_FEFCO_201
    global MNOZNIK_SZTUK_ALT_1
    global MNOZNIK_SZTUK_ALT_2
    global MNOZNIK_FEFC0_300_301
    global OPTYMALIZUJ_ODPAD
    global MIN_OSZCZEDNOSC_ODPADU_M2

    MAX_DNI_ROZNICY = int(settings["MAX_DNI_ROZNICY"])
    MAX_PRZYCINKA_PROC = float(settings["MAX_PRZYCINKA_PROC"])
    MIN_SZTUKI = float(settings["MIN_SZTUKI"])
    MIN_M2 = float(settings["MIN_M2"])
    MINIMUM_WARUNEK = str(settings["MINIMUM_WARUNEK"]).strip().upper()
    TRYB_PRZYCINKI = str(settings["TRYB_PRZYCINKI"]).strip().lower()
    LACZ_DLUGOSC_Z_SZEROKOSCIA = bool(settings["LACZ_DLUGOSC_Z_SZEROKOSCIA"])
    LACZ_ROZNE_FEFCO = bool(settings["LACZ_ROZNE_FEFCO"])
    UZYJ_ALTERNATYWNYCH_DLA_FEFCO_201 = bool(settings["UZYJ_ALTERNATYWNYCH_DLA_FEFCO_201"])
    MNOZNIK_SZTUK_ALT_1 = float(settings["MNOZNIK_SZTUK_ALT_1"])
    MNOZNIK_SZTUK_ALT_2 = float(settings["MNOZNIK_SZTUK_ALT_2"])
    MNOZNIK_FEFC0_300_301 = float(settings["MNOZNIK_FEFC0_300_301"])
    OPTYMALIZUJ_ODPAD = bool(settings["OPTYMALIZUJ_ODPAD"])
    MIN_OSZCZEDNOSC_ODPADU_M2 = float(settings["MIN_OSZCZEDNOSC_ODPADU_M2"])


def _float_from_gui(value: str, field_name: str, minimum: float = 0.0) -> float:
    text = str(value).strip().replace(" ", "").replace(",", ".")
    try:
        number = float(text)
    except Exception:
        raise ValueError(f"Pole `{field_name}` musi być liczbą.")
    if number < minimum:
        raise ValueError(f"Pole `{field_name}` nie może być mniejsze niż {minimum}.")
    return number


def _int_from_gui(value: str, field_name: str, minimum: int = 0) -> int:
    number = _float_from_gui(value, field_name, minimum)
    if abs(number - round(number)) > 0.000001:
        raise ValueError(f"Pole `{field_name}` musi być liczbą całkowitą.")
    return int(round(number))


def pobierz_ustawienia_gui() -> bool:
    """Pokazuje okno ustawień. Zwraca False, jeśli użytkownik anulował program."""
    if tk is None or ttk is None:
        zastosuj_ustawienia(wczytaj_ustawienia_z_pliku())
        return True

    settings = wczytaj_ustawienia_z_pliku()
    result = {"ok": False}

    root = tk.Tk()
    root.title("Tessainator V6 — ustawienia łączenia formatek")
    root.resizable(False, False)
    root.attributes("-topmost", True)

    main_frame = ttk.Frame(root, padding=14)
    main_frame.grid(row=0, column=0, sticky="nsew")

    ttk.Label(
        main_frame,
        text=(
            "Ustaw zasady łączenia. W wygenerowanym Excelu arkusz USTAWIENIA "
            "pokaże użyte wartości oraz opis działania każdej z nich."
        ),
        wraplength=610,
        justify="left",
    ).grid(row=0, column=0, sticky="w", pady=(0, 10))

    variables = {
        "MAX_DNI_ROZNICY": tk.StringVar(value=str(settings["MAX_DNI_ROZNICY"])),
        "MAX_PRZYCINKA_PROC": tk.StringVar(value=str(settings["MAX_PRZYCINKA_PROC"])),
        "MIN_SZTUKI": tk.StringVar(value=str(settings["MIN_SZTUKI"])),
        "MIN_M2": tk.StringVar(value=str(settings["MIN_M2"])),
        "MINIMUM_WARUNEK": tk.StringVar(value=str(settings["MINIMUM_WARUNEK"]).upper()),
        "TRYB_PRZYCINKI": tk.StringVar(value=str(settings["TRYB_PRZYCINKI"]).lower()),
        "LACZ_DLUGOSC_Z_SZEROKOSCIA": tk.BooleanVar(value=bool(settings["LACZ_DLUGOSC_Z_SZEROKOSCIA"])),
        "LACZ_ROZNE_FEFCO": tk.BooleanVar(value=bool(settings["LACZ_ROZNE_FEFCO"])),
        "UZYJ_ALTERNATYWNYCH_DLA_FEFCO_201": tk.BooleanVar(value=bool(settings["UZYJ_ALTERNATYWNYCH_DLA_FEFCO_201"])),
        "MNOZNIK_SZTUK_ALT_1": tk.StringVar(value=str(settings["MNOZNIK_SZTUK_ALT_1"])),
        "MNOZNIK_SZTUK_ALT_2": tk.StringVar(value=str(settings["MNOZNIK_SZTUK_ALT_2"])),
        "MNOZNIK_FEFC0_300_301": tk.StringVar(value=str(settings["MNOZNIK_FEFC0_300_301"])),
        "OPTYMALIZUJ_ODPAD": tk.BooleanVar(value=bool(settings["OPTYMALIZUJ_ODPAD"])),
        "MIN_OSZCZEDNOSC_ODPADU_M2": tk.StringVar(value=str(settings["MIN_OSZCZEDNOSC_ODPADU_M2"])),
    }

    notebook = ttk.Notebook(main_frame)
    notebook.grid(row=1, column=0, sticky="nsew")
    basic_frame = ttk.Frame(notebook, padding=12)
    advanced_frame = ttk.Frame(notebook, padding=12)
    notebook.add(basic_frame, text="Podstawowe")
    notebook.add(advanced_frame, text="Zaawansowane")

    def add_entry(parent, row: int, label: str, key: str, help_text: str = "") -> int:
        ttk.Label(parent, text=label).grid(row=row, column=0, sticky="w", padx=(0, 10), pady=4)
        ttk.Entry(parent, textvariable=variables[key], width=14).grid(row=row, column=1, sticky="w", pady=4)
        ttk.Label(parent, text=help_text, foreground="#555555").grid(row=row, column=2, sticky="w", pady=4)
        return row + 1

    row = 0
    row = add_entry(basic_frame, row, "Maks. różnica dat dostawy", "MAX_DNI_ROZNICY", "dni")
    row = add_entry(basic_frame, row, "Maks. przycinka", "MAX_PRZYCINKA_PROC", "%")

    ttk.Label(basic_frame, text="Sposób liczenia przycinki").grid(row=row, column=0, sticky="w", padx=(0, 10), pady=4)
    ttk.Combobox(
        basic_frame,
        textvariable=variables["TRYB_PRZYCINKI"],
        values=["area", "side", "both"],
        width=11,
        state="readonly",
    ).grid(row=row, column=1, sticky="w", pady=3)
    ttk.Label(basic_frame, text="area = pole, side = boki, both = oba warunki", foreground="#555555").grid(
        row=row, column=2, sticky="w", pady=3
    )
    row += 1

    ttk.Checkbutton(
        basic_frame,
        text="Łącz długość z szerokością i szerokość z długością (obrót o 90°)",
        variable=variables["LACZ_DLUGOSC_Z_SZEROKOSCIA"],
    ).grid(row=row, column=0, columnspan=3, sticky="w", pady=(8, 3))
    row += 1
    ttk.Label(
        basic_frame,
        text="Włącz tylko, jeśli kierunek fali, druku i wykrojnika pozwala obracać formatkę.",
        foreground="#8A4B08",
        wraplength=590,
    ).grid(row=row, column=0, columnspan=3, sticky="w", padx=(20, 0), pady=(0, 5))
    row += 1

    ttk.Checkbutton(
        basic_frame,
        text="Pozwalaj łączyć różne rodzaje FEFCO",
        variable=variables["LACZ_ROZNE_FEFCO"],
    ).grid(row=row, column=0, columnspan=3, sticky="w", pady=3)
    row += 1
    ttk.Label(
        basic_frame,
        text="Po włączeniu program może utworzyć jedną grupę z kilku kodów FEFCO.",
        foreground="#8A4B08",
    ).grid(row=row, column=0, columnspan=3, sticky="w", padx=(20, 0), pady=(0, 5))
    row += 1

    row = add_entry(basic_frame, row, "Minimum sztukowe", "MIN_SZTUKI", "szt.")
    row = add_entry(basic_frame, row, "Minimum powierzchni", "MIN_M2", "m²")

    ttk.Label(basic_frame, text="Warunek minimum").grid(row=row, column=0, sticky="w", padx=(0, 10), pady=4)
    ttk.Combobox(
        basic_frame,
        textvariable=variables["MINIMUM_WARUNEK"],
        values=["OR", "AND"],
        width=11,
        state="readonly",
    ).grid(row=row, column=1, sticky="w", pady=4)
    ttk.Label(basic_frame, text="OR = sztuki albo m², AND = oba warunki", foreground="#555555").grid(
        row=row, column=2, sticky="w", pady=4
    )

    advanced_row = 0

    ttk.Checkbutton(
        advanced_frame,
        text="Dla FEFCO 201 używaj wymiarów alternatywnych, jeśli są dostępne",
        variable=variables["UZYJ_ALTERNATYWNYCH_DLA_FEFCO_201"],
    ).grid(row=advanced_row, column=0, columnspan=3, sticky="w", pady=(0, 8))
    advanced_row += 1

    advanced_row = add_entry(advanced_frame, advanced_row, "Mnożnik jednego rodzaju alt.", "MNOZNIK_SZTUK_ALT_1", "FEFCO 201")
    advanced_row = add_entry(advanced_frame, advanced_row, "Mnożnik dwóch rodzajów alt.", "MNOZNIK_SZTUK_ALT_2", "FEFCO 201")
    advanced_row = add_entry(advanced_frame, advanced_row, "Mnożnik FEFCO 300/301", "MNOZNIK_FEFC0_300_301", "zwykle 2")

    ttk.Checkbutton(
        advanced_frame,
        text="Po utworzeniu grup dziel je, jeśli zmniejszy to łączny odpad",
        variable=variables["OPTYMALIZUJ_ODPAD"],
    ).grid(row=advanced_row, column=0, columnspan=3, sticky="w", pady=(10, 5))
    advanced_row += 1

    advanced_row = add_entry(
        advanced_frame,
        advanced_row,
        "Minimalna oszczędność odpadu",
        "MIN_OSZCZEDNOSC_ODPADU_M2",
        "m² wymagane do podziału",
    )

    ttk.Label(
        main_frame,
        text="Po zatwierdzeniu wybierzesz plik Excel. Program nie pobiera danych z bazy.",
        wraplength=610,
        justify="left",
        foreground="#555555",
    ).grid(row=2, column=0, sticky="w", pady=(10, 8))

    button_frame = ttk.Frame(main_frame)
    button_frame.grid(row=3, column=0, sticky="e", pady=(4, 0))

    def save_and_continue() -> None:
        try:
            new_settings = {
                "MAX_DNI_ROZNICY": _int_from_gui(variables["MAX_DNI_ROZNICY"].get(), "Maks. różnica dostawy", 0),
                "MAX_PRZYCINKA_PROC": _float_from_gui(variables["MAX_PRZYCINKA_PROC"].get(), "Maks. przycinka", 0),
                "MIN_SZTUKI": _float_from_gui(variables["MIN_SZTUKI"].get(), "Minimum sztukowe", 0),
                "MIN_M2": _float_from_gui(variables["MIN_M2"].get(), "Minimum metrowe", 0),
                "MINIMUM_WARUNEK": variables["MINIMUM_WARUNEK"].get().strip().upper(),
                "TRYB_PRZYCINKI": variables["TRYB_PRZYCINKI"].get().strip().lower(),
                "LACZ_DLUGOSC_Z_SZEROKOSCIA": bool(variables["LACZ_DLUGOSC_Z_SZEROKOSCIA"].get()),
                "LACZ_ROZNE_FEFCO": bool(variables["LACZ_ROZNE_FEFCO"].get()),
                "UZYJ_ALTERNATYWNYCH_DLA_FEFCO_201": bool(variables["UZYJ_ALTERNATYWNYCH_DLA_FEFCO_201"].get()),
                "MNOZNIK_SZTUK_ALT_1": _float_from_gui(variables["MNOZNIK_SZTUK_ALT_1"].get(), "Mnożnik sztuk alt. 1", 0),
                "MNOZNIK_SZTUK_ALT_2": _float_from_gui(variables["MNOZNIK_SZTUK_ALT_2"].get(), "Mnożnik sztuk alt. 2", 0),
                "MNOZNIK_FEFC0_300_301": _float_from_gui(variables["MNOZNIK_FEFC0_300_301"].get(), "Mnożnik FEFCO 300/301", 0),
                "OPTYMALIZUJ_ODPAD": bool(variables["OPTYMALIZUJ_ODPAD"].get()),
                "MIN_OSZCZEDNOSC_ODPADU_M2": _float_from_gui(
                    variables["MIN_OSZCZEDNOSC_ODPADU_M2"].get(),
                    "Minimalna oszczędność odpadu",
                    0,
                ),
            }
            if new_settings["MINIMUM_WARUNEK"] not in {"OR", "AND"}:
                raise ValueError("Warunek minimum musi mieć wartość OR albo AND.")
            if new_settings["TRYB_PRZYCINKI"] not in {"area", "side", "both"}:
                raise ValueError("Tryb przycinki musi mieć wartość area, side albo both.")
            zastosuj_ustawienia(new_settings)
            zapisz_ustawienia_do_pliku(new_settings)
            result["ok"] = True
            root.destroy()
        except Exception as exc:
            messagebox.showerror("Błąd ustawień", str(exc), parent=root)

    def cancel() -> None:
        result["ok"] = False
        root.destroy()

    ttk.Button(button_frame, text="Anuluj", command=cancel).grid(row=0, column=0, padx=(0, 8))
    ttk.Button(button_frame, text="Zapisz i wybierz plik", command=save_and_continue).grid(row=0, column=1)

    root.protocol("WM_DELETE_WINDOW", cancel)
    root.update_idletasks()
    width = root.winfo_width()
    height = root.winfo_height()
    x = (root.winfo_screenwidth() // 2) - (width // 2)
    y = (root.winfo_screenheight() // 2) - (height // 2)
    root.geometry(f"+{x}+{y}")
    root.mainloop()
    return result["ok"]


def wybierz_plik() -> Optional[Path]:
    if tk is None:
        print("Brak tkinter. Podaj ścieżkę do pliku jako pierwszy argument programu.")
        if len(sys.argv) > 1:
            return Path(sys.argv[1])
        return None

    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)

    filename = filedialog.askopenfilename(
        title="Wybierz plik Excel z zamówieniami",
        filetypes=[
            ("Excel", "*.xlsx *.xls *.xlsm"),
            ("OpenDocument Spreadsheet", "*.ods"),
            ("CSV", "*.csv"),
            ("Wszystkie pliki", "*.*"),
        ],
    )
    root.destroy()

    if not filename:
        return None
    return Path(filename)


def normalizuj_nazwy_kolumn(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df.columns = [str(c).strip().lower() for c in df.columns]
    return df


def czy_prawda(value) -> bool:
    if pd.isna(value):
        return False
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value != 0
    text = str(value).strip().lower()
    return text in {"1", "true", "t", "tak", "yes", "y", "x"}


def to_number(value) -> float:
    if pd.isna(value):
        return math.nan
    if isinstance(value, str):
        value = value.replace(" ", "").replace(",", ".")
    try:
        return float(value)
    except Exception:
        return math.nan


def normalizuj_fefco(value) -> str:
    if pd.isna(value):
        return ""
    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]
    return text


def znajdz_kolumne(df: pd.DataFrame, mozliwe: Iterable[str]) -> Optional[str]:
    cols = set(df.columns)
    for col in mozliwe:
        if col.lower() in cols:
            return col.lower()
    return None


def wybierz_kolumne_daty(df: pd.DataFrame, ostrzezenia: list[str]) -> str:
    # Aktualne wymaganie: porównujemy różnicę daty dostawy, maksymalnie 16 dni.
    # Dlatego program preferuje delivery_date z eksportu SQL.
    candidates = [
        "delivery_date",
        "data_dostawy",
        "termin_dostawy",
        "date_delivery",
        "order_created_at",
        "created_at",
        "o.created_at",
        "order_date",
        "data_zamowienia",
    ]
    col = znajdz_kolumne(df, candidates)
    if col is None:
        raise ValueError(
            "Nie znaleziono kolumny z datą dostawy. Upewnij się, że plik ma kolumnę `delivery_date` "
            "albo inną kolumnę z datą, np. `created_at/order_created_at`."
        )
    if col != "delivery_date":
        ostrzezenia.append(
            f"Uwaga: w pliku nie znaleziono `delivery_date`, więc program użył kolumny `{col}`. "
            "Jeśli chcesz liczyć różnicę dostawy, eksportuj z SQL kolumnę `i.delivery_date`."
        )
    return col


def wczytaj_plik(path: Path) -> pd.DataFrame:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return pd.read_csv(path, sep=None, engine="python")
    if suffix == ".ods":
        return pd.read_excel(path, engine="odf")
    return pd.read_excel(path)


def przygotuj_dane(df_raw: pd.DataFrame, ostrzezenia: list[str]) -> tuple[pd.DataFrame, pd.DataFrame, str]:
    df = normalizuj_nazwy_kolumn(df_raw)

    wymagane = ["fefco", "dimension_x", "dimension_y", "quantity"]
    braki = [c for c in wymagane if c not in df.columns]
    if braki:
        raise ValueError("Brakuje wymaganych kolumn: " + ", ".join(braki))

    date_col = wybierz_kolumne_daty(df, ostrzezenia)

    df["_date"] = pd.to_datetime(df[date_col], errors="coerce")
    df["_fefco"] = df["fefco"].apply(normalizuj_fefco)
    df["_base_x"] = df["dimension_x"].apply(to_number)
    df["_base_y"] = df["dimension_y"].apply(to_number)
    df["_quantity"] = df["quantity"].apply(to_number)

    effective_rows = []

    for idx, row in df.iterrows():
        fefco = row["_fefco"]
        qty_multiplier = 1
        formatka_count = 1
        source = "base"
        fefco_201_liczba_formatek = ""
        fefco_201_opis = ""
        format_variants = [
            {
                "x": row["_base_x"],
                "y": row["_base_y"],
                "source": "base",
                "copies": 1,
                "qty_multiplier": 1.0,
            }
        ]

        if UZYJ_ALTERNATYWNYCH_DLA_FEFCO_201 and fefco == "201":
            alt2_ok = (
                "has_alternative_dimensions_2" in df.columns
                and czy_prawda(row.get("has_alternative_dimensions_2"))
                and "alternative_dimension_2_x" in df.columns
                and "alternative_dimension_2_y" in df.columns
                and to_number(row.get("alternative_dimension_2_x")) > 0
                and to_number(row.get("alternative_dimension_2_y")) > 0
            )
            alt1_ok = (
                "has_alternative_dimensions" in df.columns
                and czy_prawda(row.get("has_alternative_dimensions"))
                and "alternative_dimension_x" in df.columns
                and "alternative_dimension_y" in df.columns
                and to_number(row.get("alternative_dimension_x")) > 0
                and to_number(row.get("alternative_dimension_y")) > 0
            )

            if alt1_ok and alt2_ok:
                # Dwa zestawy wymiarów to dwa różne rodzaje formatek. Łączny
                # mnożnik rozdzielamy po równo między oba rodzaje.
                multiplier_per_type = MNOZNIK_SZTUK_ALT_2 / 2
                format_variants = [
                    {
                        "x": to_number(row.get("alternative_dimension_x")),
                        "y": to_number(row.get("alternative_dimension_y")),
                        "source": "alternative_1",
                        "copies": 2,
                        "qty_multiplier": multiplier_per_type,
                    },
                    {
                        "x": to_number(row.get("alternative_dimension_2_x")),
                        "y": to_number(row.get("alternative_dimension_2_y")),
                        "source": "alternative_2",
                        "copies": 2,
                        "qty_multiplier": multiplier_per_type,
                    },
                ]
                qty_multiplier = MNOZNIK_SZTUK_ALT_2
                formatka_count = 4
                source = "alternative_1+alternative_2"
                fefco_201_liczba_formatek = 4
                fefco_201_opis = (
                    "FEFCO 201: 2 formatki alternative_1 oraz "
                    "2 formatki alternative_2"
                )
            elif alt1_ok:
                format_variants = [
                    {
                        "x": to_number(row.get("alternative_dimension_x")),
                        "y": to_number(row.get("alternative_dimension_y")),
                        "source": "alternative_1",
                        "copies": 2,
                        "qty_multiplier": MNOZNIK_SZTUK_ALT_1,
                    }
                ]
                qty_multiplier = MNOZNIK_SZTUK_ALT_1
                formatka_count = 2
                source = "alternative_1"
                fefco_201_liczba_formatek = 2
                fefco_201_opis = "FEFCO 201: pozycja składa się z 2 formatek"
            elif alt2_ok:
                # Obsługa pliku z samym drugim zestawem wymiarów: nadal jest to
                # jeden rodzaj formatki występujący dwa razy.
                format_variants = [
                    {
                        "x": to_number(row.get("alternative_dimension_2_x")),
                        "y": to_number(row.get("alternative_dimension_2_y")),
                        "source": "alternative_2",
                        "copies": 2,
                        "qty_multiplier": MNOZNIK_SZTUK_ALT_2 / 2,
                    }
                ]
                qty_multiplier = MNOZNIK_SZTUK_ALT_2 / 2
                formatka_count = 2
                source = "alternative_2"
                fefco_201_liczba_formatek = 2
                fefco_201_opis = "FEFCO 201: pozycja składa się z 2 formatek alternative_2"
            else:
                fefco_201_opis = "FEFCO 201: brak alternatywnego formatu, liczba formatek nieokreślona"

        if fefco in {"300", "301"}:
            qty_multiplier *= MNOZNIK_FEFC0_300_301
            for variant in format_variants:
                variant["qty_multiplier"] *= MNOZNIK_FEFC0_300_301

        qty = row["_quantity"]
        first_variant = format_variants[0]
        x = first_variant["x"]
        y = first_variant["y"]
        # Wartości tymczasowe służą do walidacji. Podczas rozwijania pozycji
        # zostaną policzone osobno dla każdego rodzaju formatki.
        effective_qty = (
            qty * first_variant["qty_multiplier"] / first_variant["copies"]
            if not math.isnan(qty)
            else math.nan
        )
        effective_m2_one = (x * y / 1_000_000) if not (math.isnan(x) or math.isnan(y)) else math.nan
        effective_m2_total = effective_m2_one * effective_qty if not math.isnan(effective_m2_one) and not math.isnan(effective_qty) else math.nan

        effective_rows.append(
            {
                "_effective_x": x,
                "_effective_y": y,
                "_effective_quantity": effective_qty,
                "_effective_m2_one": effective_m2_one,
                "_effective_m2_total": effective_m2_total,
                "_format_source": source,
                "_qty_multiplier": qty_multiplier,
                "_formatka_count": formatka_count,
                "_format_variants": format_variants,
                "_fefco_201_liczba_formatek": fefco_201_liczba_formatek,
                "_fefco_201_opis": fefco_201_opis,
            }
        )

    eff = pd.DataFrame(effective_rows, index=df.index)
    df = pd.concat([df, eff], axis=1)

    def powod_odrzucenia(row: pd.Series) -> str:
        reasons = []
        if pd.isna(row["_date"]):
            reasons.append("Brak lub niepoprawna data dostawy")
        if not str(row["_fefco"]).strip():
            reasons.append("Brak kodu FEFCO")
        if pd.isna(row["_effective_x"]) or float(row["_effective_x"]) <= 0:
            reasons.append("Długość formatki musi być większa od zera")
        if pd.isna(row["_effective_y"]) or float(row["_effective_y"]) <= 0:
            reasons.append("Szerokość formatki musi być większa od zera")
        if pd.isna(row["_effective_quantity"]) or float(row["_effective_quantity"]) <= 0:
            reasons.append("Ilość musi być większa od zera")
        return "; ".join(reasons)

    df["_powod_odrzucenia"] = df.apply(powod_odrzucenia, axis=1)

    invalid_mask = df["_powod_odrzucenia"] != ""

    odrzucone = df.loc[invalid_mask].copy()
    dobre = df.loc[~invalid_mask].copy()

    # Rozwijamy każdy rodzaj i każdą jego kopię na osobną formatkę. Indeks każdej
    # kopii jest unikalny, więc algorytm może przypisać je niezależnie.
    expanded_rows = []
    for source_idx, row in dobre.iterrows():
        formatka_count = int(row["_formatka_count"])
        variants = row["_format_variants"]
        for variant in variants:
            copies = int(variant["copies"])
            for formatka_nr in range(1, copies + 1):
                expanded = row.copy()
                effective_x = float(variant["x"])
                effective_y = float(variant["y"])
                effective_qty = float(row["_quantity"]) * float(variant["qty_multiplier"]) / copies
                effective_m2_one = effective_x * effective_y / 1_000_000

                expanded["_effective_x"] = effective_x
                expanded["_effective_y"] = effective_y
                expanded["_effective_quantity"] = effective_qty
                expanded["_effective_m2_one"] = effective_m2_one
                expanded["_effective_m2_total"] = effective_m2_one * effective_qty
                expanded["_format_source"] = variant["source"]
                expanded["_qty_multiplier"] = variant["qty_multiplier"]
                expanded["_source_row_index"] = source_idx
                expanded["_formatka_nr"] = formatka_nr
                expanded["_formatka_count_for_type"] = copies
                expanded_rows.append(expanded)

    if expanded_rows:
        dobre = pd.DataFrame(expanded_rows).drop(columns=["_format_variants"]).reset_index(drop=True)
    else:
        dobre = dobre.copy()
        dobre["_source_row_index"] = pd.Series(dtype="object")
        dobre["_formatka_nr"] = pd.Series(dtype="int64")
        dobre["_formatka_count_for_type"] = pd.Series(dtype="int64")
        dobre = dobre.drop(columns=["_format_variants"], errors="ignore")

    odrzucone = odrzucone.drop(columns=["_format_variants"], errors="ignore")

    if len(odrzucone) > 0:
        ostrzezenia.append(f"Odrzucono {len(odrzucone)} wierszy z powodu braku/niepoprawnej daty, FEFCO, wymiarów lub ilości.")

    return dobre, odrzucone, date_col


def trim_for_dims(part_x: float, part_y: float, target_x: float, target_y: float) -> Optional[dict]:
    """Sprawdza przycinkę przy zachowaniu oryginalnej orientacji X × Y."""
    if part_x <= 0 or part_y <= 0 or target_x <= 0 or target_y <= 0:
        return None
    if part_x > target_x or part_y > target_y:
        return None

    trim_x_pct = ((target_x - part_x) / target_x) * 100
    trim_y_pct = ((target_y - part_y) / target_y) * 100
    area_trim_pct = ((target_x * target_y - part_x * part_y) / (target_x * target_y)) * 100
    max_side_trim_pct = max(trim_x_pct, trim_y_pct)

    if TRYB_PRZYCINKI == "area":
        ok = area_trim_pct <= MAX_PRZYCINKA_PROC
    elif TRYB_PRZYCINKI == "side":
        ok = max_side_trim_pct <= MAX_PRZYCINKA_PROC
    elif TRYB_PRZYCINKI == "both":
        ok = area_trim_pct <= MAX_PRZYCINKA_PROC and max_side_trim_pct <= MAX_PRZYCINKA_PROC
    else:
        raise ValueError("TRYB_PRZYCINKI musi być: area, side albo both")

    if not ok:
        return None
    return {
        "trim_area_pct": area_trim_pct,
        "trim_x_pct": trim_x_pct,
        "trim_y_pct": trim_y_pct,
        "max_side_trim_pct": max_side_trim_pct,
    }


def best_target_for_indices(df: pd.DataFrame, indices: list[int]) -> TargetCheck:
    if not indices:
        return TargetCheck(False, reason="pusta grupa")

    x_sides: list[float] = []
    y_sides: list[float] = []
    for idx in indices:
        x = float(df.at[idx, "_effective_x"])
        y = float(df.at[idx, "_effective_y"])
        x_sides.append(x)
        y_sides.append(y)

    if not LACZ_DLUGOSC_Z_SZEROKOSCIA:
        # Zachowujemy orientację każdej formatki: X porównujemy tylko z X, a Y z Y.
        candidate_targets = [(max(x_sides), max(y_sides))]
    else:
        # Przy dozwolonym obrocie każda krawędź może stać się X albo Y. Wspólny
        # format zawsze ma boki równe co najmniej jednemu z boków wejściowych,
        # więc wystarczy sprawdzić wszystkie takie pary. Dłuższy bok pokazujemy
        # jako X, aby wynik był jednoznaczny.
        sides = sorted(set(x_sides + y_sides))
        candidate_targets = [
            (target_x, target_y)
            for target_x in sides
            for target_y in sides
            if target_x >= target_y
        ]

    best_result: Optional[TargetCheck] = None
    best_key = None
    for target_x, target_y in candidate_targets:
        trims = []
        rotated_indices: set[int] = set()
        valid = True
        for idx in indices:
            px = float(df.at[idx, "_effective_x"])
            py = float(df.at[idx, "_effective_y"])
            checks = [(False, trim_for_dims(px, py, target_x, target_y))]
            if LACZ_DLUGOSC_Z_SZEROKOSCIA and px != py:
                checks.append((True, trim_for_dims(py, px, target_x, target_y)))
            checks = [(rotated, check) for rotated, check in checks if check is not None]
            if not checks:
                valid = False
                break
            metric_name = "max_side_trim_pct" if TRYB_PRZYCINKI == "side" else "trim_area_pct"
            rotated, best_check = min(
                checks,
                key=lambda item: (item[1][metric_name], item[0]),
            )
            if rotated:
                rotated_indices.add(idx)
            trims.append(best_check[metric_name])

        if not valid or not trims:
            continue

        result = TargetCheck(
            valid=True,
            target_x=target_x,
            target_y=target_y,
            max_trim_pct=max(trims),
            avg_trim_pct=sum(trims) / len(trims),
            reason="OK",
            rotated_indices=frozenset(rotated_indices),
        )
        # Najpierw najmniejszy zamawiany arkusz, potem najmniejsza przycinka.
        key = (target_x * target_y, result.max_trim_pct, result.avg_trim_pct, target_x, target_y)
        if best_key is None or key < best_key:
            best_key = key
            best_result = result

    if best_result is not None:
        return best_result
    return TargetCheck(False, reason=f"brak wspólnego formatu przy max {MAX_PRZYCINKA_PROC}%")


def iteruj_zbiory_do_laczenia(df: pd.DataFrame):
    """Zwraca zbiory kandydatów zgodnie z ustawieniem mieszania kodów FEFCO."""
    if LACZ_ROZNE_FEFCO:
        yield "RÓŻNE FEFCO", df
        return
    yield from df.groupby("_fefco", sort=True)


def opis_fefco_dla_indeksow(df: pd.DataFrame, indices: list[int]) -> str:
    values = sorted({str(value) for value in df.loc[indices, "_fefco"] if str(value).strip()})
    return " + ".join(values)


def date_window_ok(df: pd.DataFrame, indices: list[int]) -> bool:
    dates = df.loc[indices, "_date"]
    return (dates.max() - dates.min()).days <= MAX_DNI_ROZNICY


def totals_for_indices(df: pd.DataFrame, indices: list[int]) -> dict:
    qty = float(df.loc[indices, "_effective_quantity"].sum())
    m2 = float(df.loc[indices, "_effective_m2_total"].sum())
    return {"qty": qty, "m2": m2}


def fmt_mm(value: float) -> str:
    if math.isnan(value):
        return ""
    if abs(value - round(value)) < 0.0001:
        return str(int(round(value)))
    return f"{value:.2f}".rstrip("0").rstrip(".")


def opis_formatu_zamawianego(target_x: float, target_y: float) -> str:
    return f"{fmt_mm(target_x)} x {fmt_mm(target_y)} mm"


def waste_metrics_for_indices(df: pd.DataFrame, indices: list[int], target: TargetCheck) -> dict:
    """Liczy odpad względem zamawianego arkusza w formacie jednej formatki."""
    ordered_m2_one = (target.target_x * target.target_y) / 1_000_000
    total_ordered_m2 = 0.0
    total_waste_m2 = 0.0
    rows: dict[int, dict] = {}

    for idx in indices:
        part_x = float(df.at[idx, "_effective_x"])
        part_y = float(df.at[idx, "_effective_y"])
        qty = float(df.at[idx, "_effective_quantity"])
        part_m2_one = (part_x * part_y) / 1_000_000
        waste_m2_one = max(ordered_m2_one - part_m2_one, 0.0)
        waste_pct_one = (waste_m2_one / ordered_m2_one * 100) if ordered_m2_one > 0 else math.nan
        ordered_m2_total = ordered_m2_one * qty
        waste_m2_total = waste_m2_one * qty

        rows[idx] = {
            "m2_zamowione_1_arkusz": ordered_m2_one,
            "m2_zamowione_razem": ordered_m2_total,
            "odpad_m2_na_arkusz": waste_m2_one,
            "odpad_proc_arkusza": waste_pct_one,
            "odpad_m2_razem": waste_m2_total,
        }
        total_ordered_m2 += ordered_m2_total
        total_waste_m2 += waste_m2_total

    total_waste_pct = (total_waste_m2 / total_ordered_m2 * 100) if total_ordered_m2 > 0 else math.nan
    return {
        "rows": rows,
        "m2_zamowione_razem": total_ordered_m2,
        "odpad_m2_razem": total_waste_m2,
        "odpad_proc_razem": total_waste_pct,
    }


def minimum_ok(qty: float, m2: float) -> bool:
    mode = MINIMUM_WARUNEK.strip().upper()
    if mode == "OR":
        return qty >= MIN_SZTUKI or m2 >= MIN_M2
    if mode == "AND":
        return qty >= MIN_SZTUKI and m2 >= MIN_M2
    raise ValueError("MINIMUM_WARUNEK musi być: OR albo AND")


def _ocen_grupe(df: pd.DataFrame, indices: list[int]) -> Optional[dict]:
    """Zwraca pełną ocenę poprawnej grupy albo None."""
    # Podczas pierwszego przebiegu program nadal tworzy grupy z co najmniej
    # dwóch pozycji. Optymalizacja może jednak wydzielić jedną formatkę,
    # jeżeli sama spełnia minimum sztukowe/metrowe. Zamówienie jej w jej
    # własnym formacie daje wtedy zerową przycinkę i może znacznie obniżyć
    # odpad pozostałej grupy.
    if not indices or not date_window_ok(df, indices):
        return None

    totals = totals_for_indices(df, indices)
    if not minimum_ok(totals["qty"], totals["m2"]):
        return None

    target = best_target_for_indices(df, indices)
    if not target.valid:
        return None

    waste = waste_metrics_for_indices(df, indices, target)
    return {
        "indices": list(indices),
        "target": target,
        "totals": totals,
        "waste": waste,
    }


def _kandydaci_podzialu(df: pd.DataFrame, indices: list[int]):
    """Tworzy ograniczony zestaw sensownych podziałów według wymiarów."""
    if len(indices) < 2:
        return

    def dimensions(idx: int) -> tuple[float, float, float, float]:
        x = float(df.at[idx, "_effective_x"])
        y = float(df.at[idx, "_effective_y"])
        short = min(x, y)
        long = max(x, y)
        area = short * long
        ratio = long / short if short > 0 else math.inf
        return short, long, area, ratio

    dimensions_by_index = {idx: dimensions(idx) for idx in indices}
    orders = [
        sorted(indices, key=lambda idx: (dimensions_by_index[idx][0], dimensions_by_index[idx][1], idx)),
        sorted(indices, key=lambda idx: (dimensions_by_index[idx][1], dimensions_by_index[idx][0], idx)),
        sorted(indices, key=lambda idx: (dimensions_by_index[idx][2], dimensions_by_index[idx][0], idx)),
        sorted(indices, key=lambda idx: (dimensions_by_index[idx][3], dimensions_by_index[idx][2], idx)),
    ]

    seen = set()

    # Program stworzony przez Maksymiliana Dylę, 2026.
    # Każdą formatkę sprawdzamy jako potencjalne samodzielne zamówienie,
    # nie tylko te leżące na krańcach sortowania wymiarów. Dzięki temu
    # optymalizacja nie zależy od tego, czy kandydat jest najmniejszy/największy.
    for idx in indices:
        left = [idx]
        right = [other for other in indices if other != idx]
        partition_key = frozenset((frozenset(left), frozenset(right)))
        if partition_key in seen:
            continue
        seen.add(partition_key)
        yield left, right

    for ordered in orders:
        # Dopuszczamy część jednoelementową. Jej poprawność zostanie
        # niżej zweryfikowana tymi samymi progami MIN_SZTUKI i MIN_M2.
        for cut in range(1, len(ordered)):
            left = ordered[:cut]
            right = ordered[cut:]
            partition_key = frozenset((frozenset(left), frozenset(right)))
            if partition_key in seen:
                continue
            seen.add(partition_key)
            yield left, right


def optymalizuj_grupe_pod_katem_odpadu(df: pd.DataFrame, indices: list[int]) -> list[list[int]]:
    """
    Rekurencyjnie dzieli grupę tylko wtedy, gdy obie części nadal spełniają
    wszystkie warunki, a ich łączny odpad jest mniejszy.
    """
    evaluation_cache: dict[tuple[int, ...], Optional[dict]] = {}

    def evaluate(group: list[int]) -> Optional[dict]:
        key = tuple(sorted(group))
        if key not in evaluation_cache:
            evaluation_cache[key] = _ocen_grupe(df, list(key))
        return evaluation_cache[key]

    def optimize(group: list[int]) -> list[list[int]]:
        baseline = evaluate(group)
        if baseline is None or len(group) < 2:
            return [group]

        baseline_waste = float(baseline["waste"]["odpad_m2_razem"])
        best_split = None
        best_key = None

        for left, right in _kandydaci_podzialu(df, group):
            left_eval = evaluate(left)
            right_eval = evaluate(right)
            if left_eval is None or right_eval is None:
                continue

            split_waste = (
                float(left_eval["waste"]["odpad_m2_razem"])
                + float(right_eval["waste"]["odpad_m2_razem"])
            )
            saving = baseline_waste - split_waste
            if saving <= max(MIN_OSZCZEDNOSC_ODPADU_M2, 0.0) + 1e-9:
                continue

            # Najpierw największa oszczędność, potem mniejszy łączny odpad
            # i bardziej równy liczebnie podział.
            key = (-saving, split_waste, abs(len(left) - len(right)))
            if best_key is None or key < best_key:
                best_key = key
                best_split = (left, right)

        if best_split is None:
            return [group]

        left, right = best_split
        return optimize(left) + optimize(right)

    return optimize(list(indices))


def id_pozycji(row: pd.Series, idx: int) -> str:
    order_id = row.get("order_id", "")
    position = row.get("position", "")
    item_name = row.get("item_name", "")
    parts = []
    if not pd.isna(order_id) and str(order_id).strip() != "":
        parts.append(f"order:{order_id}")
    if not pd.isna(position) and str(position).strip() != "":
        parts.append(f"poz:{position}")
    if not pd.isna(item_name) and str(item_name).strip() != "":
        parts.append(str(item_name))
    if not parts:
        parts.append(f"wiersz:{idx + 2}")
    formatka_count = int(row.get("_formatka_count", 1))
    formatka_count_for_type = int(row.get("_formatka_count_for_type", 1))
    formatka_nr = int(row.get("_formatka_nr", 1))
    if formatka_count > 1:
        source = str(row.get("_format_source", "formatka"))
        parts.append(f"{source}:{formatka_nr}/{formatka_count_for_type}")
    return " / ".join(parts)


def znajdz_dopasowania_par(df: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for _fefco_group, part in iteruj_zbiory_do_laczenia(df):
        idxs = list(part.index)
        for a, b in combinations(idxs, 2):
            if not date_window_ok(df, [a, b]):
                continue
            target = best_target_for_indices(df, [a, b])
            if not target.valid:
                continue
            totals = totals_for_indices(df, [a, b])
            if not minimum_ok(totals["qty"], totals["m2"]):
                continue

            waste = waste_metrics_for_indices(df, [a, b], target)
            date_min = df.loc[[a, b], "_date"].min()
            date_max = df.loc[[a, b], "_date"].max()
            rows.append(
                {
                    "FEFCO": opis_fefco_dla_indeksow(df, [a, b]),
                    "Order ID 1": df.at[a, "order_id"] if "order_id" in df.columns else "",
                    "Pozycja 1": df.at[a, "position"] if "position" in df.columns else "",
                    "Produkt 1": df.at[a, "item_name"] if "item_name" in df.columns else "",
                    "Formatka 1": f"{df.at[a, '_format_source']} {df.at[a, '_formatka_nr']}",
                    "Order ID 2": df.at[b, "order_id"] if "order_id" in df.columns else "",
                    "Pozycja 2": df.at[b, "position"] if "position" in df.columns else "",
                    "Produkt 2": df.at[b, "item_name"] if "item_name" in df.columns else "",
                    "Formatka 2": f"{df.at[b, '_format_source']} {df.at[b, '_formatka_nr']}",
                    "Najwcześniejsza dostawa": date_min.date(),
                    "Najpóźniejsza dostawa": date_max.date(),
                    "Różnica dat [dni]": (date_max - date_min).days,
                    "Wspólny format [mm]": opis_formatu_zamawianego(target.target_x, target.target_y),
                    "Maks. przycinka [%]": round(target.max_trim_pct, 2),
                    "Odpad [%]": round(waste["odpad_proc_razem"], 2),
                    "Odpad [m²]": round(waste["odpad_m2_razem"], 4),
                    "Powierzchnia zamówiona [m²]": round(waste["m2_zamowione_razem"], 4),
                    "Sztuki razem": round(totals["qty"], 2),
                    "Powierzchnia wykorzystana [m²]": round(totals["m2"], 4),
                }
            )
    return pd.DataFrame(rows)


def znajdz_grupy_greedy(df: pd.DataFrame) -> tuple[pd.DataFrame, dict[int, str], dict[str, TargetCheck]]:
    assigned: dict[int, str] = {}
    group_targets: dict[str, TargetCheck] = {}
    group_rows = []
    optimization_rows = []
    group_no = 1
    optimization_no = 1

    for _fefco_group, part in iteruj_zbiory_do_laczenia(df.sort_values("_date")):
        indices = list(part.index)
        for start_idx in indices:
            if start_idx in assigned:
                continue

            group = [start_idx]
            changed = True
            while changed:
                changed = False
                best_candidate = None
                best_candidate_target = None
                best_key = None

                for cand_idx in indices:
                    if cand_idx in assigned or cand_idx in group:
                        continue
                    candidate_group = group + [cand_idx]
                    if not date_window_ok(df, candidate_group):
                        continue
                    target = best_target_for_indices(df, candidate_group)
                    if not target.valid:
                        continue

                    totals = totals_for_indices(df, candidate_group)
                    # Kandydat wybierany tak, żeby grupa była jak najbardziej spójna formatowo.
                    key = (target.max_trim_pct, abs(totals["qty"] - MIN_SZTUKI), abs(totals["m2"] - MIN_M2))
                    if best_key is None or key < best_key:
                        best_key = key
                        best_candidate = cand_idx
                        best_candidate_target = target

                if best_candidate is not None:
                    group.append(best_candidate)
                    changed = True

            if len(group) < 2:
                continue

            totals = totals_for_indices(df, group)
            if not minimum_ok(totals["qty"], totals["m2"]):
                continue

            target = best_target_for_indices(df, group)
            if not target.valid:
                continue

            waste_before = waste_metrics_for_indices(df, group, target)
            optimized_groups = (
                optymalizuj_grupe_pod_katem_odpadu(df, group)
                if OPTYMALIZUJ_ODPAD
                else [group]
            )
            optimized_evaluations = [_ocen_grupe(df, optimized_group) for optimized_group in optimized_groups]
            optimized_evaluations = [evaluation for evaluation in optimized_evaluations if evaluation is not None]
            waste_after = sum(
                float(evaluation["waste"]["odpad_m2_razem"])
                for evaluation in optimized_evaluations
            )
            saving = max(float(waste_before["odpad_m2_razem"]) - waste_after, 0.0)
            optimization_id = f"O{optimization_no}"
            optimization_no += 1

            optimization_rows.append(
                {
                    "ID optymalizacji": optimization_id,
                    "fefco": opis_fefco_dla_indeksow(df, group),
                    "liczba_grup_przed": 1,
                    "liczba_grup_po": len(optimized_evaluations),
                    "odpad_m2_przed": round(float(waste_before["odpad_m2_razem"]), 4),
                    "odpad_m2_po": round(waste_after, 4),
                    "oszczednosc_odpadu_m2": round(saving, 4),
                    "oszczednosc_odpadu_proc": round(
                        saving / float(waste_before["odpad_m2_razem"]) * 100,
                        2,
                    )
                    if float(waste_before["odpad_m2_razem"]) > 0
                    else 0.0,
                    "decyzja": "PODZIELONO" if len(optimized_evaluations) > 1 else "BEZ ZMIAN",
                    "format_przed": opis_formatu_zamawianego(target.target_x, target.target_y),
                    "formaty_po": " | ".join(
                        opis_formatu_zamawianego(
                            evaluation["target"].target_x,
                            evaluation["target"].target_y,
                        )
                        for evaluation in optimized_evaluations
                    ),
                }
            )

            for evaluation in optimized_evaluations:
                optimized_group = evaluation["indices"]
                optimized_target = evaluation["target"]
                optimized_totals = evaluation["totals"]
                optimized_waste = evaluation["waste"]
                group_id = f"F{group_no}"
                group_no += 1

                for idx in optimized_group:
                    assigned[idx] = group_id
                group_targets[group_id] = optimized_target

                date_min = df.loc[optimized_group, "_date"].min()
                date_max = df.loc[optimized_group, "_date"].max()
                group_rows.append(
                    {
                        "ID grupy": group_id,
                        "wynik_optymalizacji": (
                            "PODZIELONO DLA MNIEJSZEGO ODPADU"
                            if len(optimized_evaluations) > 1
                            else "BEZ ZMIAN"
                        ),
                        "format_zamawiany": opis_formatu_zamawianego(
                            optimized_target.target_x,
                            optimized_target.target_y,
                        ),
                        "Sztuki do zamówienia": round(optimized_totals["qty"], 2),
                        "fefco": opis_fefco_dla_indeksow(df, optimized_group),
                        "liczba_zamowien": int(df.loc[optimized_group, "order_id"].nunique()) if "order_id" in df.columns else "",
                        "liczba_pozycji": int(df.loc[optimized_group, "_source_row_index"].nunique()),
                        "liczba_formatek": len(optimized_group),
                        "data_od": date_min.date(),
                        "data_do": date_max.date(),
                        "roznica_dni": (date_max - date_min).days,
                        "format_zamawiany_info": "Zamówić format wielkości jednej formatki",
                        "format_docelowy_x": round(optimized_target.target_x, 2),
                        "format_docelowy_y": round(optimized_target.target_y, 2),
                        "max_przycinka_proc": round(optimized_target.max_trim_pct, 2),
                        "srednia_przycinka_proc": round(optimized_target.avg_trim_pct, 2),
                        "odpad_proc_razem": round(optimized_waste["odpad_proc_razem"], 2),
                        "odpad_m2_razem": round(optimized_waste["odpad_m2_razem"], 4),
                        "m2_zamowione_razem": round(optimized_waste["m2_zamowione_razem"], 4),
                        "m2_razem_efektywne": round(optimized_totals["m2"], 4),
                        "pozycje": " | ".join(id_pozycji(df.loc[i], i) for i in optimized_group),
                        "ID optymalizacji": optimization_id,
                    }
                )

    groups_df = pd.DataFrame(group_rows)
    groups_df.attrs["optymalizacja"] = optimization_rows
    return groups_df, assigned, group_targets


def przygotuj_grupy_operatora(grupy: pd.DataFrame) -> pd.DataFrame:
    """Buduje zwarty arkusz z informacjami potrzebnymi do złożenia zamówienia."""
    if grupy.empty:
        return grupy.copy()
    out = grupy.copy()
    out["Decyzja"] = out["wynik_optymalizacji"].map(
        {
            "PODZIELONO DLA MNIEJSZEGO ODPADU": "PODZIELONO — MNIEJSZY ODPAD",
            "BEZ ZMIAN": "ZAMÓW JAK WYLICZONO",
        }
    ).fillna(out["wynik_optymalizacji"])
    out = out.rename(
        columns={
            "format_zamawiany": "Format do zamówienia [mm]",
            "Sztuki do zamówienia": "Liczba sztuk do zamówienia",
            "fefco": "FEFCO",
            "liczba_zamowien": "Liczba zamówień",
            "liczba_pozycji": "Liczba pozycji",
            "liczba_formatek": "Liczba formatek",
            "data_od": "Najwcześniejsza dostawa",
            "data_do": "Najpóźniejsza dostawa",
            "roznica_dni": "Różnica dat [dni]",
            "max_przycinka_proc": "Maks. przycinka [%]",
            "srednia_przycinka_proc": "Średnia przycinka [%]",
            "odpad_proc_razem": "Odpad [%]",
            "odpad_m2_razem": "Odpad [m²]",
            "m2_zamowione_razem": "Powierzchnia zamówiona [m²]",
            "m2_razem_efektywne": "Powierzchnia wykorzystana [m²]",
            "pozycje": "Pozycje należące do grupy",
        }
    )
    columns = [
        "ID grupy", "Decyzja", "FEFCO", "Format do zamówienia [mm]",
        "Liczba sztuk do zamówienia", "Najwcześniejsza dostawa",
        "Najpóźniejsza dostawa", "Różnica dat [dni]", "Liczba zamówień",
        "Liczba pozycji", "Liczba formatek", "Maks. przycinka [%]",
        "Średnia przycinka [%]", "Odpad [%]", "Odpad [m²]",
        "Powierzchnia zamówiona [m²]", "Powierzchnia wykorzystana [m²]",
        "ID optymalizacji", "Pozycje należące do grupy",
    ]
    return out[[column for column in columns if column in out.columns]]


def przygotuj_pozycje_output(df: pd.DataFrame, assigned: dict[int, str], group_targets: dict[str, TargetCheck]) -> pd.DataFrame:
    out = df.copy()

    # Najważniejsze kolumny na początku arkusza POZYCJE.
    group_values = [assigned.get(idx, "") for idx in out.index]
    order_values = out.pop("order_id") if "order_id" in out.columns else pd.Series("", index=out.index)
    position_values = out.pop("position") if "position" in out.columns else pd.Series("", index=out.index)
    formatka_source_values = out.pop("_format_source")
    formatka_nr_values = out.pop("_formatka_nr")
    formatka_count_for_type_values = out.pop("_formatka_count_for_type")
    formatka_count_values = out.pop("_formatka_count")
    format_kartonu_values = [
        opis_formatu_zamawianego(float(row["_effective_x"]), float(row["_effective_y"]))
        for _, row in out.iterrows()
    ]
    status_values = ["POŁĄCZONO" if idx in assigned else "NIEPOŁĄCZONE" for idx in out.index]

    out.insert(0, "ID grupy", group_values)
    out.insert(1, "Status", status_values)
    out.insert(2, "Order ID", order_values)
    out.insert(3, "Pozycja", position_values)
    out.insert(4, "Format wejściowy [mm]", format_kartonu_values)
    out.insert(5, "Rodzaj formatki", formatka_source_values)
    out.insert(6, "Numer formatki", formatka_nr_values)
    out.insert(7, "Liczba formatek danego rodzaju", formatka_count_for_type_values)
    out.insert(8, "Łączna liczba formatek w pozycji", formatka_count_values)

    target_format = []
    target_x = []
    target_y = []
    rotated_values = []
    trim_pct = []
    waste_pct = []
    waste_m2_one = []
    waste_m2_total = []
    ordered_m2_one = []
    ordered_m2_total = []

    for idx, row in out.iterrows():
        gid = assigned.get(idx)
        if not gid:
            target_format.append("")
            target_x.append("")
            target_y.append("")
            rotated_values.append("NIE")
            trim_pct.append("")
            waste_pct.append("")
            waste_m2_one.append("")
            waste_m2_total.append("")
            ordered_m2_one.append("")
            ordered_m2_total.append("")
            continue

        target = group_targets[gid]
        target_format.append(opis_formatu_zamawianego(target.target_x, target.target_y))
        target_x.append(round(target.target_x, 2))
        target_y.append(round(target.target_y, 2))

        rotated = idx in target.rotated_indices
        rotated_values.append("TAK" if rotated else "NIE")
        part_x = float(row["_effective_y"] if rotated else row["_effective_x"])
        part_y = float(row["_effective_x"] if rotated else row["_effective_y"])
        check = trim_for_dims(part_x, part_y, target.target_x, target.target_y)
        if check:
            trim_pct.append(round(check["trim_area_pct"] if TRYB_PRZYCINKI != "side" else check["max_side_trim_pct"], 2))
        else:
            trim_pct.append("")

        waste = waste_metrics_for_indices(out, [idx], target)["rows"][idx]
        waste_pct.append(round(waste["odpad_proc_arkusza"], 2))
        waste_m2_one.append(round(waste["odpad_m2_na_arkusz"], 4))
        waste_m2_total.append(round(waste["odpad_m2_razem"], 4))
        ordered_m2_one.append(round(waste["m2_zamowione_1_arkusz"], 4))
        ordered_m2_total.append(round(waste["m2_zamowione_razem"], 4))

    out.insert(9, "Format do zamówienia [mm]", target_format)
    out.insert(10, "Obrócono o 90°", rotated_values)
    out.insert(11, "Przycinka [%]", trim_pct)
    out.insert(12, "Odpad [%]", waste_pct)
    out.insert(13, "Odpad na arkusz [m²]", waste_m2_one)
    out.insert(14, "Odpad łącznie [m²]", waste_m2_total)
    out.insert(15, "Powierzchnia zamówiona na arkusz [m²]", ordered_m2_one)
    out.insert(16, "Powierzchnia zamówiona łącznie [m²]", ordered_m2_total)
    out.insert(17, "format_docelowy_x", target_x)
    out.insert(18, "format_docelowy_y", target_y)

    # Czytelne kolumny techniczne.
    rename_map = {
        "_date": "data_uzyta_do_porownania",
        "_fefco": "fefco_norm",
        "_effective_x": "format_x_uzyty",
        "_effective_y": "format_y_uzyty",
        "_effective_quantity": "Sztuki efektywne",
        "_effective_m2_one": "m2_1szt_efektywne",
        "_effective_m2_total": "Powierzchnia wykorzystana [m²]",
        "_format_source": "zrodlo_formatu",
        "_qty_multiplier": "mnoznik_sztuk",
        "_source_row_index": "indeks_pozycji_zrodlowej",
        "_fefco_201_liczba_formatek": "fefco_201_liczba_formatek",
        "_fefco_201_opis": "fefco_201_opis",
        "_powod_odrzucenia": "Powód odrzucenia",
        "client_order_number": "Numer zamówienia klienta",
        "client_name": "Klient",
        "item_name": "Nazwa produktu",
        "delivery_date": "Data dostawy",
        "fefco": "FEFCO",
    }
    out = out.rename(columns=rename_map)
    if "Data dostawy" in out.columns:
        out["Data dostawy"] = pd.to_datetime(out["Data dostawy"], errors="coerce").dt.date
    primary_columns = [
        "ID grupy", "Status", "Order ID", "Pozycja", "Numer zamówienia klienta",
        "Klient", "Nazwa produktu", "FEFCO", "Data dostawy", "Rodzaj formatki",
        "Numer formatki", "Łączna liczba formatek w pozycji", "Format wejściowy [mm]",
        "Format do zamówienia [mm]", "Obrócono o 90°", "Sztuki efektywne",
        "Przycinka [%]", "Odpad [%]", "Odpad łącznie [m²]",
        "Powierzchnia zamówiona łącznie [m²]", "Powierzchnia wykorzystana [m²]",
    ]
    existing_primary = [column for column in primary_columns if column in out.columns]
    remaining = [column for column in out.columns if column not in existing_primary]
    out = out[existing_primary + remaining]
    status_order = pd.Categorical(out["Status"], ["POŁĄCZONO", "NIEPOŁĄCZONE"], ordered=True)
    out = out.assign(_kolejnosc_statusu=status_order).sort_values(
        ["_kolejnosc_statusu", "ID grupy", "Order ID", "Pozycja", "Rodzaj formatki", "Numer formatki"],
        na_position="last",
    ).drop(columns=["_kolejnosc_statusu"]).reset_index(drop=True)
    return out


def przygotuj_pozycje_operatora(pozycje_techniczne: pd.DataFrame) -> pd.DataFrame:
    """Ogranicza widok operatora do kolumn potrzebnych do podjęcia decyzji."""
    columns = [
        "ID grupy", "Status", "Order ID", "Pozycja", "Numer zamówienia klienta",
        "Klient", "Nazwa produktu", "FEFCO", "Data dostawy", "Rodzaj formatki",
        "Numer formatki", "Łączna liczba formatek w pozycji", "Format wejściowy [mm]",
        "Format do zamówienia [mm]", "Obrócono o 90°", "Sztuki efektywne",
        "Przycinka [%]", "Odpad [%]", "Odpad łącznie [m²]",
        "Powierzchnia zamówiona łącznie [m²]", "Powierzchnia wykorzystana [m²]",
    ]
    return pozycje_techniczne[[column for column in columns if column in pozycje_techniczne.columns]].copy()


def przygotuj_odrzucone_output(odrzucone: pd.DataFrame) -> pd.DataFrame:
    if odrzucone.empty:
        return odrzucone.copy()
    out = odrzucone.rename(
        columns={
            "_powod_odrzucenia": "Powód odrzucenia",
            "order_id": "Order ID",
            "position": "Pozycja",
            "client_name": "Klient",
            "item_name": "Nazwa produktu",
            "fefco": "FEFCO",
            "dimension_x": "Długość [mm]",
            "dimension_y": "Szerokość [mm]",
            "quantity": "Ilość [szt.]",
            "delivery_date": "Data dostawy",
        }
    )
    columns = [
        "Powód odrzucenia", "Order ID", "Pozycja", "Klient", "Nazwa produktu",
        "FEFCO", "Długość [mm]", "Szerokość [mm]", "Ilość [szt.]", "Data dostawy",
    ]
    return out[[column for column in columns if column in out.columns]].copy()


def przygotuj_ostrzezenia_output(ostrzezenia: list[str]) -> pd.DataFrame:
    if not ostrzezenia:
        return pd.DataFrame(
            [["INFORMACJA", "Program", "Brak ostrzeżeń.", "Nie jest wymagane żadne działanie."]],
            columns=["Poziom", "Obszar", "Komunikat", "Zalecane działanie"],
        )
    rows = []
    for warning in ostrzezenia:
        area = "Dane wejściowe"
        action = "Sprawdź wskazane wiersze i popraw plik źródłowy."
        if "delivery_date" in warning or "kolumny" in warning.lower():
            area = "Kolumny wejściowe"
            action = "Sprawdź nazwy kolumn oraz źródło daty."
        rows.append(["OSTRZEŻENIE", area, warning, action])
    return pd.DataFrame(rows, columns=["Poziom", "Obszar", "Komunikat", "Zalecane działanie"])


def przygotuj_optymalizacje_output(optymalizacja: pd.DataFrame) -> pd.DataFrame:
    if optymalizacja.empty:
        return optymalizacja.copy()
    out = optymalizacja.rename(
        columns={
            "decyzja": "Decyzja",
            "fefco": "FEFCO",
            "format_przed": "Format przed [mm]",
            "formaty_po": "Formaty po podziale [mm]",
            "liczba_grup_przed": "Liczba grup przed",
            "liczba_grup_po": "Liczba grup po",
            "odpad_m2_przed": "Odpad przed [m²]",
            "odpad_m2_po": "Odpad po [m²]",
            "oszczednosc_odpadu_m2": "Oszczędność [m²]",
            "oszczednosc_odpadu_proc": "Oszczędność [%]",
        }
    )
    columns = [
        "ID optymalizacji", "Decyzja", "FEFCO", "Format przed [mm]",
        "Formaty po podziale [mm]", "Liczba grup przed", "Liczba grup po",
        "Odpad przed [m²]", "Odpad po [m²]", "Oszczędność [m²]", "Oszczędność [%]",
    ]
    return out[[column for column in columns if column in out.columns]].copy()


def ustaw_szerokosci_excel(writer: pd.ExcelWriter) -> None:
    """Nadaje arkuszom V6 spójny układ operatora i formaty danych."""
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    cienka_szara = Side(style="thin", color="B7C9D6")
    lekka_krawedz = Border(
        left=cienka_szara,
        right=cienka_szara,
        top=cienka_szara,
        bottom=cienka_szara,
    )
    naglowek_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    wiersz_parzysty_fill = PatternFill(fill_type="solid", fgColor="EAF3F8")
    wiersz_nieparzysty_fill = PatternFill(fill_type="solid", fgColor="FFFFFF")
    green_fill = PatternFill(fill_type="solid", fgColor="E2F0D9")
    amber_fill = PatternFill(fill_type="solid", fgColor="FFF2CC")
    red_fill = PatternFill(fill_type="solid", fgColor="FCE4D6")

    tab_colors = {
        "PODSUMOWANIE": "1F4E78",
        "GRUPY DO ZAMÓWIENIA": "2F75B5",
        "POZYCJE": "5B9BD5",
        "ODRZUCONE": "C00000",
        "OSTRZEŻENIA": "FFC000",
        "USTAWIENIA": "70AD47",
        "OPTYMALIZACJA": "A5A5A5",
    }

    for sheet_name, worksheet in writer.sheets.items():
        worksheet.sheet_view.showGridLines = False
        if sheet_name in tab_colors:
            worksheet.sheet_properties.tabColor = tab_colors[sheet_name]

        for column_cells in worksheet.columns:
            max_len = 0
            column_letter = column_cells[0].column_letter
            header = str(column_cells[0].value or "")
            for cell in column_cells:
                try:
                    value = "" if cell.value is None else str(cell.value)
                    max_len = max(max_len, len(value))
                except Exception:
                    pass
            max_width = 72 if header in {"Zasada działania", "Pozycje należące do grupy", "Komunikat", "Zalecane działanie"} else 45
            worksheet.column_dimensions[column_letter].width = min(max(max_len + 2, 10), max_width)

        for row_number, row_cells in enumerate(worksheet.iter_rows(), start=1):
            if row_number == 1:
                fill = naglowek_fill
            elif row_number % 2 == 0:
                fill = wiersz_parzysty_fill
            else:
                fill = wiersz_nieparzysty_fill

            for cell in row_cells:
                cell.border = lekka_krawedz
                cell.fill = fill
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if row_number == 1:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        freeze_by_sheet = {
            "GRUPY DO ZAMÓWIENIA": "D2",
            "POZYCJE": "E2",
            "DANE TECHNICZNE": "E2",
            "DOPASOWANIA TECHNICZNE": "D2",
        }
        worksheet.freeze_panes = freeze_by_sheet.get(sheet_name, "A2")
        if sheet_name not in {"PODSUMOWANIE", "OSTRZEŻENIA", "USTAWIENIA"}:
            worksheet.auto_filter.ref = worksheet.dimensions

        headers = {cell.value: cell.column for cell in worksheet[1]}
        for header, column_number in headers.items():
            header_text = str(header or "")
            for row_number in range(2, worksheet.max_row + 1):
                cell = worksheet.cell(row=row_number, column=column_number)
                if "dostawa" in header_text.lower() or header_text == "Data dostawy":
                    cell.number_format = "dd.mm.yyyy"
                elif "[%]" in header_text:
                    cell.number_format = '0.00" %"'
                elif "[m²]" in header_text:
                    cell.number_format = "#,##0.0000"
                elif any(word in header_text for word in ["Liczba", "Sztuki", "Ilość", "Różnica dat"]):
                    cell.number_format = "#,##0.##"

        status_column = headers.get("Status")
        if status_column:
            for row_number in range(2, worksheet.max_row + 1):
                cell = worksheet.cell(row=row_number, column=status_column)
                cell.fill = green_fill if cell.value == "POŁĄCZONO" else amber_fill
                cell.font = Font(bold=True, color="006100" if cell.value == "POŁĄCZONO" else "9C6500")

        decision_column = headers.get("Decyzja")
        if decision_column:
            for row_number in range(2, worksheet.max_row + 1):
                cell = worksheet.cell(row=row_number, column=decision_column)
                cell.fill = green_fill if "ZAMÓW" in str(cell.value) else amber_fill
                cell.font = Font(bold=True)

        level_column = headers.get("Poziom")
        if level_column:
            for row_number in range(2, worksheet.max_row + 1):
                cell = worksheet.cell(row=row_number, column=level_column)
                cell.fill = red_fill if cell.value == "BŁĄD" else amber_fill if cell.value == "OSTRZEŻENIE" else green_fill
                cell.font = Font(bold=True)

        if sheet_name == "USTAWIENIA" and "Wartość" in headers and "Zasada działania" in headers:
            value_column = headers["Wartość"]
            explanation_column = headers["Zasada działania"]
            worksheet["A1"].comment = Comment(
                "Arkusz dokumentuje dokładne ustawienia użyte do utworzenia tego wyniku.",
                "Tessainator V6",
            )
            for row_number in range(2, worksheet.max_row + 1):
                explanation = worksheet.cell(row=row_number, column=explanation_column).value
                if explanation:
                    worksheet.cell(row=row_number, column=value_column).comment = Comment(str(explanation), "Tessainator V6")

        if sheet_name == "GRUPY DO ZAMÓWIENIA":
            for row_number in range(2, worksheet.max_row + 1):
                worksheet.row_dimensions[row_number].height = 45


def _excel_sum_formula(df: pd.DataFrame, column_name: str, sheet_name: str) -> str:
    if column_name not in df.columns or df.empty:
        return "=0"
    from openpyxl.utils import get_column_letter

    column_letter = get_column_letter(df.columns.get_loc(column_name) + 1)
    return f"=SUM('{sheet_name}'!${column_letter}$2:${column_letter}${len(df) + 1})"


def przygotuj_podsumowanie(
    pozycje: pd.DataFrame,
    optymalizacja: pd.DataFrame,
    grupy: Optional[pd.DataFrame] = None,
    odrzucone: Optional[pd.DataFrame] = None,
) -> pd.DataFrame:
    """Buduje formułowe podsumowanie wszystkich powierzchni i odpadu."""
    from openpyxl.utils import get_column_letter

    effective_all = _excel_sum_formula(pozycje, "Powierzchnia wykorzystana [m²]", "POZYCJE")
    ordered = _excel_sum_formula(pozycje, "Powierzchnia zamówiona łącznie [m²]", "POZYCJE")
    waste = _excel_sum_formula(pozycje, "Odpad łącznie [m²]", "POZYCJE")

    if {"Status", "Powierzchnia wykorzystana [m²]"}.issubset(pozycje.columns) and not pozycje.empty:
        status_col = get_column_letter(pozycje.columns.get_loc("Status") + 1)
        effective_col = get_column_letter(pozycje.columns.get_loc("Powierzchnia wykorzystana [m²]") + 1)
        last_row = len(pozycje) + 1
        effective_joined = (
            f'=SUMIF(\'POZYCJE\'!${status_col}$2:${status_col}${last_row},'
            f'"POŁĄCZONO",\'POZYCJE\'!${effective_col}$2:${effective_col}${last_row})'
        )
    else:
        effective_joined = "=0"

    saving = _excel_sum_formula(
        optymalizacja,
        "Oszczędność [m²]",
        "OPTYMALIZACJA",
    )
    group_count = 0 if grupy is None or grupy.empty else len(grupy)
    joined_count = int((pozycje.get("Status", pd.Series(dtype=str)) == "POŁĄCZONO").sum())
    unjoined_count = int((pozycje.get("Status", pd.Series(dtype=str)) == "NIEPOŁĄCZONE").sum())
    rejected_count = 0 if odrzucone is None or odrzucone.empty else len(odrzucone)

    return pd.DataFrame(
        [
            ["Liczba utworzonych grup", group_count, "Grupy gotowe do złożenia wspólnego zamówienia."],
            ["Liczba połączonych formatek", joined_count, "Formatki przypisane do gotowych grup."],
            ["Liczba niepołączonych formatek", unjoined_count, "Formatki, które nie spełniły warunków grupowania."],
            ["Liczba odrzuconych pozycji", rejected_count, "Wiersze z brakującymi lub niepoprawnymi danymi."],
            [
                "SUMA M² WSZYSTKICH FORMATEK",
                effective_all,
                "Wszystkie rodzaje formatek z arkusza POZYCJE, także niepołączone.",
            ],
            [
                "M² efektywne w gotowych grupach",
                effective_joined,
                "Rzeczywista powierzchnia formatek przypisanych do grup.",
            ],
            [
                "M² zamówione dla gotowych grup",
                ordered,
                "Powierzchnia arkuszy zamawianych po dobraniu wspólnego formatu.",
            ],
            [
                "Odpad m² po optymalizacji",
                waste,
                "Różnica między powierzchnią zamawianą a efektywną.",
            ],
            [
                "Efektywność wykorzystania materiału",
                "=IF(B8=0,0,B7/B8)",
                "M² efektywne w grupach podzielone przez m² zamówione.",
            ],
            [
                "Oszczędność odpadu dzięki podziałom",
                saving,
                "Spadek odpadu względem grup utworzonych przed drugim przebiegiem.",
            ],
            [
                "Kontrola bilansu m²",
                "=B8-B7-B9",
                "Powinna być bliska 0; różnice mogą wynikać wyłącznie z zaokrągleń.",
            ],
        ],
        columns=["Miara", "Wartość", "Opis"],
    )


def ustaw_format_podsumowania(writer: pd.ExcelWriter) -> None:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    worksheet = writer.sheets["PODSUMOWANIE"]
    worksheet.sheet_view.showGridLines = False
    worksheet.column_dimensions["A"].width = 40
    worksheet.column_dimensions["B"].width = 20
    worksheet.column_dimensions["C"].width = 72

    worksheet["A2"].fill = PatternFill(fill_type="solid", fgColor="70AD47")
    worksheet["B2"].fill = PatternFill(fill_type="solid", fgColor="70AD47")
    worksheet["C2"].fill = PatternFill(fill_type="solid", fgColor="E2F0D9")
    worksheet["A2"].font = Font(bold=True, color="FFFFFF")
    worksheet["B2"].font = Font(bold=True, color="FFFFFF", size=14)

    for row in range(2, worksheet.max_row + 1):
        worksheet[f"B{row}"].number_format = "#,##0.0000"
        worksheet[f"B{row}"].alignment = Alignment(horizontal="right", vertical="center")
        worksheet.row_dimensions[row].height = 28
        metric = worksheet[f"A{row}"].value
        if isinstance(metric, str) and metric.startswith("Liczba"):
            worksheet[f"B{row}"].number_format = "#,##0"
        elif metric == "Efektywność wykorzystania materiału":
            worksheet[f"B{row}"].number_format = "0.00%"

    accent = Side(style="medium", color="1F4E78")
    for cell in worksheet[2]:
        cell.border = Border(top=accent, bottom=accent)


def _awaryjna_sciezka_wyniku(output_path: Path) -> Path:
    """Zwraca wolną nazwę dla wyniku, którego nie można zastąpić."""
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    base_name = f"{output_path.stem}_{timestamp}"
    candidate = output_path.with_name(base_name + output_path.suffix)
    counter = 2

    while candidate.exists():
        candidate = output_path.with_name(f"{base_name}_{counter}{output_path.suffix}")
        counter += 1

    return candidate


def zapisz_wynik(
    input_path: Path,
    grupy: pd.DataFrame,
    dopasowania: pd.DataFrame,
    pozycje: pd.DataFrame,
    odrzucone: pd.DataFrame,
    ostrzezenia: list[str],
    date_col: str,
) -> Path:
    output_path = input_path.with_name(input_path.stem + "_polaczone_formatki.xlsx")
    optimization_rows = list(grupy.attrs.get("optymalizacja", []))
    optymalizacja = przygotuj_optymalizacje_output(pd.DataFrame(optimization_rows))
    grupy_operatora = przygotuj_grupy_operatora(grupy)
    pozycje_operatora = przygotuj_pozycje_operatora(pozycje)
    odrzucone_operatora = przygotuj_odrzucone_output(odrzucone)
    ostrz = przygotuj_ostrzezenia_output(ostrzezenia)
    podsumowanie = przygotuj_podsumowanie(
        pozycje_operatora,
        optymalizacja,
        grupy_operatora,
        odrzucone_operatora,
    )

    ustawienia = pd.DataFrame(
        [
            ["Daty i zgodność", "Maksymalna różnica dat dostawy", MAX_DNI_ROZNICY, "Pozycje mogą trafić do jednej grupy, jeśli skrajne daty różnią się najwyżej o tę liczbę dni.", "MAX_DNI_ROZNICY"],
            ["Wymiary i obrót", "Maksymalna przycinka", MAX_PRZYCINKA_PROC, "Odrzuca wspólny format, gdy wymagana przycinka przekracza podany procent.", "MAX_PRZYCINKA_PROC"],
            ["Wymiary i obrót", "Sposób liczenia przycinki", TRYB_PRZYCINKI, "area = ubytek pola; side = największa przycinka boku; both = oba warunki muszą być spełnione.", "TRYB_PRZYCINKI"],
            ["Wymiary i obrót", "Łączenie długości z szerokością", "TAK" if LACZ_DLUGOSC_Z_SZEROKOSCIA else "NIE", "TAK pozwala obrócić formatkę o 90°. Należy sprawdzić kierunek fali, druku i wykrojnika.", "LACZ_DLUGOSC_Z_SZEROKOSCIA"],
            ["FEFCO", "Łączenie różnych kodów FEFCO", "TAK" if LACZ_ROZNE_FEFCO else "NIE", "TAK pozwala tworzyć wspólną grupę z różnych kodów FEFCO; zgodność technologiczną trzeba potwierdzić.", "LACZ_ROZNE_FEFCO"],
            ["Minimalna partia", "Minimum sztukowe", MIN_SZTUKI, "Minimalna efektywna liczba sztuk wymagana przez warunek minimum.", "MIN_SZTUKI"],
            ["Minimalna partia", "Minimum powierzchni", MIN_M2, "Minimalna efektywna powierzchnia grupy w m².", "MIN_M2"],
            ["Minimalna partia", "Sposób spełnienia minimum", MINIMUM_WARUNEK, "OR = wystarczy minimum sztuk albo m²; AND = grupa musi spełnić oba minima.", "MINIMUM_WARUNEK"],
            ["Źródło danych", "Kolumna daty użyta do porównania", date_col, "Ta kolumna z pliku wejściowego została użyta jako data dostawy.", "KOLUMNA_DATY_UŻYTA"],
            ["FEFCO", "Alternatywne wymiary FEFCO 201", "TAK" if UZYJ_ALTERNATYWNYCH_DLA_FEFCO_201 else "NIE", "TAK rozwija dostępne alternatywne wymiary FEFCO 201 na osobne formatki.", "ALT_FEFCO_201"],
            ["FEFCO", "Mnożnik jednego rodzaju alternatywnego", MNOZNIK_SZTUK_ALT_1, "Łączny mnożnik ilości, gdy FEFCO 201 ma jeden rodzaj formatki alternatywnej.", "MNOZNIK_ALT_1"],
            ["FEFCO", "Mnożnik dwóch rodzajów alternatywnych", MNOZNIK_SZTUK_ALT_2, "Łączny mnożnik ilości, gdy FEFCO 201 ma dwa rodzaje formatek alternatywnych.", "MNOZNIK_ALT_2"],
            ["FEFCO", "Mnożnik FEFCO 300/301", MNOZNIK_FEFC0_300_301, "Mnoży efektywną ilość i powierzchnię pozycji FEFCO 300 oraz 301.", "MNOZNIK_FEFCO_300_301"],
            ["Optymalizacja", "Dzielenie grup dla mniejszego odpadu", "TAK" if OPTYMALIZUJ_ODPAD else "NIE", "TAK pozwala podzielić poprawną grupę, jeśli każda część nadal spełnia warunki i zmniejsza odpad.", "OPTYMALIZUJ_ODPAD"],
            ["Optymalizacja", "Minimalna oszczędność odpadu", MIN_OSZCZEDNOSC_ODPADU_M2, "Podział następuje tylko wtedy, gdy zmniejszy odpad co najmniej o tę liczbę m².", "MIN_OSZCZEDNOSC_ODPADU_M2"],
        ],
        columns=["Kategoria", "Ustawienie", "Wartość", "Zasada działania", "Klucz techniczny"],
    )

    if grupy_operatora.empty:
        grupy_operatora = pd.DataFrame({"Informacja": ["Nie znaleziono grup spełniających ustawione warunki."]})
    if dopasowania.empty:
        dopasowania = pd.DataFrame({"Informacja": ["Nie znaleziono par spełniających ustawione warunki."]})
    if odrzucone_operatora.empty:
        odrzucone_operatora = pd.DataFrame({"Informacja": ["Brak odrzuconych wierszy."]})
    if optymalizacja.empty:
        optymalizacja = pd.DataFrame({"Informacja": ["Brak grup wymagających analizy odpadu."]})

    # Najpierw tworzymy kompletny plik tymczasowy. Dzięki temu przerwany zapis
    # nie uszkodzi poprzedniego wyniku.
    with tempfile.NamedTemporaryFile(
        prefix=f".{output_path.stem}_",
        suffix=".xlsx",
        dir=output_path.parent,
        delete=False,
    ) as temporary_file:
        temporary_path = Path(temporary_file.name)

    try:
        with pd.ExcelWriter(temporary_path, engine="openpyxl") as writer:
            podsumowanie.to_excel(writer, index=False, sheet_name="PODSUMOWANIE")
            grupy_operatora.to_excel(writer, index=False, sheet_name="GRUPY DO ZAMÓWIENIA")
            pozycje_operatora.to_excel(writer, index=False, sheet_name="POZYCJE")
            odrzucone_operatora.to_excel(writer, index=False, sheet_name="ODRZUCONE")
            ostrz.to_excel(writer, index=False, sheet_name="OSTRZEŻENIA")
            ustawienia.to_excel(writer, index=False, sheet_name="USTAWIENIA")
            optymalizacja.to_excel(writer, index=False, sheet_name="OPTYMALIZACJA")
            dopasowania.to_excel(writer, index=False, sheet_name="DOPASOWANIA TECHNICZNE")
            pozycje.to_excel(writer, index=False, sheet_name="DANE TECHNICZNE")
            ustaw_szerokosci_excel(writer)
            ustaw_format_podsumowania(writer)
            writer.book.calculation.calcMode = "auto"
            writer.book.calculation.fullCalcOnLoad = True
            writer.book.calculation.forceFullCalc = True

        try:
            temporary_path.replace(output_path)
            return output_path
        except PermissionError:
            # Excel blokuje zastąpienie otwartego skoroszytu. Zachowujemy nowy
            # wynik pod nazwą z datą, zamiast kończyć cały program błędem.
            fallback_path = _awaryjna_sciezka_wyniku(output_path)
            temporary_path.replace(fallback_path)
            return fallback_path
    finally:
        if temporary_path.exists():
            try:
                temporary_path.unlink()
            except OSError:
                pass


def main() -> None:
    try:
        if not pobierz_ustawienia_gui():
            return

        input_path = wybierz_plik()
        if input_path is None:
            return
        if not input_path.exists():
            raise FileNotFoundError(f"Nie znaleziono pliku: {input_path}")

        ostrzezenia: list[str] = []
        df_raw = wczytaj_plik(input_path)
        df, odrzucone, date_col = przygotuj_dane(df_raw, ostrzezenia)

        if df.empty:
            raise ValueError("Po odrzuceniu błędnych wierszy nie zostały żadne poprawne pozycje do porównania.")

        dopasowania = znajdz_dopasowania_par(df)
        grupy, assigned, group_targets = znajdz_grupy_greedy(df)
        pozycje = przygotuj_pozycje_output(df, assigned, group_targets)

        expected_output_path = input_path.with_name(input_path.stem + "_polaczone_formatki.xlsx")
        output_path = zapisz_wynik(input_path, grupy, dopasowania, pozycje, odrzucone, ostrzezenia, date_col)
        fallback_note = ""
        if output_path != expected_output_path:
            fallback_note = (
                "\n\nPoprzedniego wyniku nie można było zastąpić "
                "(najczęściej dlatego, że jest otwarty w Excelu). "
                "Nowy wynik zapisano pod nazwą z datą i godziną."
            )

        show_info(
            "Gotowe",
            "Zakończono łączenie formatek.\n\n"
            f"Zapisano plik:\n{output_path}\n\n"
            "Zacznij od arkuszy: PODSUMOWANIE, GRUPY DO ZAMÓWIENIA i POZYCJE. "
            "Szczegółowe zasady użyte w tym przebiegu znajdziesz w arkuszu USTAWIENIA."
            f"{fallback_note}",
        )

    except Exception as exc:
        details = traceback.format_exc()
        show_error("Błąd programu", f"{exc}\n\nSzczegóły:\n{details}")
        raise


if __name__ == "__main__":
    main()
# Program stworzony przez Maksymiliana Dylę, 2026.